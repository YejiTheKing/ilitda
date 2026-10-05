import zipfile

import docx
import openpyxl
import pytest

from ilitda.runlog import RunLog
from ilitda.safety.mask import mask
from ilitda.safety.paths import NotAllowedError, check_allowed
from ilitda.tools.ingest import ingest
from ilitda.tools.parse import UnreadableError, parse_document
from ilitda.tools.scan import scan_folder


@pytest.fixture
def folder(tmp_path):
    d = tmp_path / "자료"
    d.mkdir()
    (d / "절차메모.txt").write_text("1. 접수\n2. 검토\n3. 발송", encoding="utf-8")
    (d / "절차메모 - 복사본.txt").write_text("1. 접수\n2. 검토\n3. 발송", encoding="utf-8")

    doc = docx.Document()
    doc.add_paragraph("견적서 작성 절차")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "유효기간 30일"
    doc.save(d / "견적서_양식.docx")

    wb = openpyxl.Workbook()
    wb.active.title = "체크리스트"
    wb.active.append(["단계", "확인"])
    wb.active.append(["사업자등록증 수령", "필수"])
    wb.save(d / "체크리스트.xlsx")

    with zipfile.ZipFile(d / "안내문.hwpx", "w") as z:
        z.writestr("Contents/section0.xml",
                   '<s xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
                   "<hp:p><hp:run><hp:t>제출 기한은 매월 10일</hp:t></hp:run></hp:p></s>")

    (d / "옛양식.hwp").write_bytes(b"\xd0\xcf\x11\xe0binary")
    (d / "~$견적서_양식.docx").write_bytes(b"lock")
    return d


def test_scan_skips_lock_files_and_flags_unreadable(folder, tmp_path):
    files = scan_folder(folder, allowed=[tmp_path])
    names = [f["name"] for f in files]
    assert "~$견적서_양식.docx" not in names
    # .hwp는 글자를 읽을 수 있지만, 가짜 바이트로 만든 이 파일은 구조가 아니라 읽기에서 거절된다
    assert {f["name"]: f["readable"] for f in files}["옛양식.hwp"] is True


def test_parse_each_format(folder, tmp_path):
    allowed = [tmp_path]
    assert "유효기간 30일" in parse_document(folder / "견적서_양식.docx", allowed)["text"]
    assert "사업자등록증 수령" in parse_document(folder / "체크리스트.xlsx", allowed)["text"]
    assert "매월 10일" in parse_document(folder / "안내문.hwpx", allowed)["text"]
    with pytest.raises(UnreadableError):
        parse_document(folder / "옛양식.hwp", allowed)


def test_ingest_continues_past_unreadable_and_logs(folder, tmp_path):
    log = RunLog(root=tmp_path / "runs")
    result = ingest(folder, log, allowed=[tmp_path])

    assert len(result["documents"]) == 4
    assert [u["name"] for u in result["unreadable"]] == ["옛양식.hwp"]
    assert len(result["copies"]) == 1

    entries = log.read()
    assert entries[0]["tool"] == "scan_folder"
    assert [e["seq"] for e in entries] == list(range(1, len(entries) + 1))
    assert any(e["kind"] == "feedback" for e in entries)


def test_folder_outside_allowed_is_refused(folder, tmp_path):
    inside = tmp_path / "허용"
    inside.mkdir()
    with pytest.raises(NotAllowedError):
        scan_folder(folder, allowed=[inside])
    with pytest.raises(NotAllowedError):
        check_allowed(inside / ".." / "자료", allowed=[inside])


def test_mask_hides_personal_number_formats():
    text = ("주민 900101-1234567, 사업자 123-45-67890, 연락처 010-1234-5678, 메일 a@b.co.kr, "
            "계좌 110-234-567890, 차량 12가 3456, 주소 부산 강서구 미음산단로 123, 기한 2026-10-06")
    masked, counts = mask(text)
    for secret in ("900101-1234567", "123-45-67890", "010-1234-5678", "a@b.co.kr", "110-234-567890", "12가 3456", "미음산단로 123"):
        assert secret not in masked
    # 날짜는 버전 판별에 쓰이므로 남아 있어야 한다
    assert "2026-10-06" in masked
    assert counts == {"주민등록번호": 1, "사업자등록번호": 1, "전화번호": 1, "계좌번호": 1, "이메일": 1, "주소": 1, "차량번호": 1}
