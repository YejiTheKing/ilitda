"""인수인계서 Word 생성: 절차 표와 '확인할 것'(~하기 꼴)만 있고, 출처·체크리스트·물음표는 없는지."""
import io

from docx import Document

from ilitda.ui.handover_doc import build_docx, file_name

STATE = {
    "goal": {"task": "신규 거래처 등록부터 첫 견적서 발송까지", "newcomer": "박신입", "owner": "이대표", "done_when": "혼자 끝까지 수행"},
    "steps": [
        {"order": 1, "action": "사업자등록증 사본을 받는다.", "who": "영업", "form": "체크리스트.xlsx", "status": "ok",
         "evidence": [{"doc": "체크리스트.xlsx", "quote": "1 | 사업자등록증 사본 수령", "found": True}]},
        {"order": 2, "action": "근거가 없는 단계", "who": "", "form": "", "status": "unverified",
         "evidence": [{"doc": "메모.txt", "quote": "없는 문장", "found": False}]},
        {"order": 3, "action": "책임자가 추가한 단계", "who": "경리", "form": "", "status": "ok", "edited_by": "이대표",
         "evidence": [{"doc": "책임자 확인", "quote": "책임자가 추가한 단계", "found": True}]},
    ],
    "answers": [
        {"topic": "승인 방법", "question": "승인은?", "answer": "결재 시스템", "source": "owner"},
        {"topic": "단가표", "question": "단가표는?", "answer": "", "source": "unresolved"},
    ],
    "versions": [{"latest": "양식_수정.docx", "family": ["양식.docx", "양식_수정.docx"], "evidence": ["내용 속 가장 최근 날짜 2025.03"], "confident": True}],
    "changes": [{"type": "기준값", "old": "유효기간: 30일", "new": "유효기간: 14일", "impact": "14일로 기재", "matters": True}],
    "checklist": [{"order": 1, "text": "사본 받기", "how_to_verify": "파일 확인", "required": True}],
    "progress": {"1": {"done": True}},
    "documents": [{"name": "체크리스트.xlsx"}],
    "unreadable": [{"name": "단가표.hwp", "reason": "손상"}],
    "status": "completed",
    "approval": {"approved": True, "note": "확인"},
}


def _text(data):
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_checklist_rows_take_who_and_form_from_closest_step():
    text = _text(build_docx(STATE, "x"))
    assert "사본 받기" in text and "영업 | 체크리스트.xlsx" not in text   # 표 셀은 따로 들어간다
    doc = Document(io.BytesIO(build_docx(STATE, "x")))
    row = doc.tables[0].rows[1].cells
    assert [c.text for c in row][:2] == ["1", "사본 받기"] and row[2].text == "영업" and row[3].text == "체크리스트.xlsx"
    held = {**STATE, "checklist": STATE["checklist"] + [{"order": 2, "text": "[미확정] 단가표 — 책임자도 확인하지 못함", "required": False}]}
    table = Document(io.BytesIO(build_docx(held, "x"))).tables[0]
    assert len(table.rows) == 2 and "[미확정]" not in _text(build_docx(held, "x"))   # 미확정 항목은 표에 없다


def test_build_docx_contains_sections():
    STATE_NO_CL = {k: v for k, v in STATE.items() if k != "checklist"}
    data = build_docx(STATE_NO_CL, "20261005-000000-abcdef")
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "신규 거래처 등록부터 첫 견적서 발송까지 인수인계서" in text
    assert "사업자등록증 사본을 받는다." in text and "체크리스트.xlsx" in text
    assert "확인할 것" in text and "단가표 확인하기" in text and "2번 '근거가 없는 단계' 책임자에게 확인하기" in text
    assert "3번" not in text.split("확인할 것")[1]   # 책임자가 추가한 단계는 확인 대상이 아니다
    for banned in ("출처", "체크리스트\n", "확인 필요", "유효기간: 14일", "구버전", "실행 ID", "?"):
        assert banned not in text, banned


def test_file_name_is_safe():
    assert file_name(STATE, "20261005-000000-abcdef").endswith("_20261005-000000-abcdef.docx")
    assert "/" not in file_name({"goal": {"task": "a/b:c"}}, "x")
