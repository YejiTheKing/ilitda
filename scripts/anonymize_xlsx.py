"""엑셀 인수인계표를 열 구조를 살려 가명 처리한다.

'업체명' 열의 값은 협력사01, 02…로, '담당자·책임자·GM·기타' 열의 값은 담당자01, 02…(직함 유지)로
일관되게 바꾸고, 같은 이름이 자유 텍스트에 나와도 같은 가명으로 바꾼다. 전화·주소·이메일·차량번호와
파일 경로(X:\\..., \\\\서버\\...)는 가린다. 시트 이름의 사람 이름도 바꾼다.

사용: .venv/bin/python scripts/anonymize_xlsx.py data/real/업무인수인계/업무인수인계.xlsx
결과: data/real/업무인수인계_비식별/업무인수인계_비식별.xlsx + 치환표.json (실명→가명, 로컬 보관)
"""
import json
import os
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT / "src"))

from ilitda.safety.mask import mask  # noqa: E402

COMPANY_HEADERS = {"업체명", "협력업체", "거래처", "회사명", "협력사"}
PERSON_HEADERS = {"담당자", "책임자", "GM", "담당", "작성자", "승인자"}
# '기타'는 '책임자' 머리글 바로 아래 칸(GM | 기타)일 때만 사람 열이다. '항목' 열의 값으로도 쓰이는 낱말이라서
SUB_HEADERS = {"기타"}
TITLES = "사장|이사|과장|대리|차장|부장|주임|전무|상무|공장장|사원|팀장|기사|반장|조장|실장|소장|대표"
NAME_TITLE = re.compile(rf"(?<![가-힣])([가-힣]{{2,4}})\s?({TITLES})(?![가-힣])")
# 셀 안의 파일 경로는 시작부터 셀 끝까지 가린다 (폴더 이름에 고객·제품 코드가 섞여 있어서)
PATH = re.compile(r"(?:[A-Za-z]:\\|\\\\)[^\n]*")
# 사람 이름으로 보려면 흔한 성씨로 시작해야 한다. 부서·일반 명사가 이름으로 잡히는 것을 막는다
SURNAMES = set("김이박최정강조윤장임한오서신권황안송류전홍고문양손배백허유남심노하곽성차주우구민나진지엄채원천방공현함변염여추도소석선설마길연위표명기반왕금옥육인맹제모탁국어은편용예경봉사부가복태목형피두감승")
NOT_NAMES = {"생산", "품질", "영업", "구매", "설비", "환경", "안전", "기술", "관리", "총무", "경리", "공무", "자재", "개발", "연구",
             "품보", "생기", "보전", "고객사", "협력사", "업체", "기타", "담당", "담당자", "책임자", "본사", "공장", "현장", "외주", "사내"}
GENERIC_COMPANY = {"고객사", "협력사", "업체", "기타", "-", "사내", "사내 시스템", "자체", "본사", "없음"}


def looks_like_person(name):
    return 2 <= len(name) <= 4 and name[0] in SURNAMES and name not in NOT_NAMES
CUSTOMERS = {"NISSAN": "고객사A", "닛산": "고객사A"}


class Mapper:
    def __init__(self):
        self.company, self.person = {}, {}

    def comp(self, name):
        key = re.sub(r"\s+", "", name)
        if key not in self.company:
            self.company[key] = f"협력사{len(self.company) + 1:02d}"
        return self.company[key]

    def pers(self, name):
        key = re.sub(r"\s+", "", name)
        if key not in self.person:
            self.person[key] = f"담당자{len(self.person) + 1:02d}"
        return self.person[key]


def main(path):
    src = Path(path)
    out_dir = src.parent.parent / (src.parent.name + "_비식별")
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(str(src))
    m = Mapper()
    masked_total = {}

    # 1) 열 구조로 회사·사람 이름을 먼저 모은다 (일관된 가명을 위해)
    layout = {}
    for ws in wb.worksheets:
        comp_cols, pers_cols = set(), set()
        for row in ws.iter_rows():
            for c in row:
                v = str(c.value).strip() if c.value is not None else ""
                if v in COMPANY_HEADERS:
                    comp_cols.add(c.column)
                elif v in PERSON_HEADERS:
                    pers_cols.add(c.column)
                elif v in SUB_HEADERS and c.row > 1:
                    above = [ws.cell(row=c.row - 1, column=col).value for col in (c.column - 1, c.column) if col >= 1]
                    if any(str(a).strip() == "책임자" for a in above if a is not None):
                        pers_cols.add(c.column)
        layout[ws.title] = (comp_cols, pers_cols)
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str) or not c.value.strip():
                    continue
                v = c.value.strip()
                if c.column in comp_cols and v not in COMPANY_HEADERS and v not in GENERIC_COMPANY and len(v) <= 20:
                    m.comp(v)
                elif c.column in pers_cols and v not in PERSON_HEADERS:
                    nm = NAME_TITLE.match(v)
                    cand = nm.group(1) if nm else v.split()[0]
                    if looks_like_person(cand):
                        m.pers(cand)
    # 시트 이름의 사람 이름
    for ws in wb.worksheets:
        first = ws.title.split("_")[0]
        if re.fullmatch(r"[가-힣]{2,4}", first) and looks_like_person(first):
            m.pers(first)
    # 자유 텍스트의 '이름 직함'
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    for nm, _ in NAME_TITLE.findall(c.value):
                        if looks_like_person(nm):
                            m.pers(nm)

    # 2) 모든 셀에 적용
    comp_items = sorted(m.company.items(), key=lambda kv: -len(kv[0]))
    pers_items = sorted(m.person.items(), key=lambda kv: -len(kv[0]))

    def scrub(text):
        for real, fake in comp_items:
            text = re.sub(re.escape(real).replace(r"\ ", r"\s*"), fake, text)
            # 띄어 쓴 표기도 잡는다 ("강일 공업" 등)
            spaced = r"\s*".join(map(re.escape, real))
            text = re.sub(spaced, fake, text)
        for real, fake in pers_items:
            text = re.sub(r"\s*".join(map(re.escape, real)), fake, text)
        for real, fake in CUSTOMERS.items():
            text = text.replace(real, fake)
        text, n = PATH.subn("[경로 가림]", text)
        if n:
            masked_total["경로"] = masked_total.get("경로", 0) + n
        text, counts = mask(text)
        for k, v in counts.items():
            masked_total[k] = masked_total.get(k, 0) + v
        return text

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip():
                    c.value = scrub(c.value)
        new_title = ws.title
        for real, fake in pers_items:
            new_title = new_title.replace(real, fake)
        ws.title = new_title[:31]

    out = out_dir / (src.stem + "_비식별.xlsx")
    wb.save(str(out))
    table = {**{k: {"kind": "회사", "replace_with": v} for k, v in m.company.items()},
             **{k: {"kind": "사람", "replace_with": v} for k, v in m.person.items()},
             **{k: {"kind": "고객사", "replace_with": v} for k, v in CUSTOMERS.items()}}
    (out_dir / "치환표.json").write_text(json.dumps(table, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{out.name}: 회사 {len(m.company)}곳, 사람 {len(m.person)}명 가명 처리, 가림 {masked_total}")
    print(f"치환표: {out_dir / '치환표.json'} (실명이 들어 있으니 저장소에 올리지 않음)")

    # 3) 검증: 결과물에서 실명이 남았는지
    from ilitda.tools.parse import parse_document
    text = parse_document(out, allowed=[out_dir.parent])["text"]
    leftovers = [k for k in list(m.company) + list(m.person) if k in re.sub(r"\s+", "", text)]
    phones = re.findall(r"0\d{1,2}[-.\s]\d{3,4}[-.\s]\d{4}", text)
    print(f"검증: 남은 실명 {len(leftovers)}개 {leftovers[:5]}, 남은 전화번호 {len(phones)}개")
    return 0 if not leftovers and not phones else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
