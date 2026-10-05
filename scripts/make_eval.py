"""변경점 탐지 평가용 양식 쌍을 만든다 (시드 평가).

- data/eval/seeded/changed/: 구버전·최신본 사이에 항목 20개를 바꾼 쌍
- data/eval/seeded/unchanged/: 내용이 같은 쌍 (오탐 측정용)
- data/eval/seeded/gold.json: 정답 (바꾼 항목 이름과 종류)

이 세트는 개발자가 만든 시드 평가다. 제3자 블라인드 평가는 scripts/evaluate_changes.py의
--detect-only 모드로 탐지 목록만 뽑아 출제자가 채점한다.

사용: .venv/bin/python scripts/make_eval.py
"""
import json
import random
from pathlib import Path

import docx

OUT = Path("data/eval/seeded")

BASE_FIELDS = [
    ("문서번호", "PO-____-___"), ("수신", "______ 귀중"), ("발주일자", "20__년 __월 __일"),
    ("납기", "발주일로부터 20일"), ("납품장소", "당사 1공장 자재창고"), ("결제조건", "납품 후 45일 현금"),
    ("검수기준", "납품 후 3일 이내 수입검사"), ("불합격 처리", "반품 및 재납품"), ("포장", "파렛트 단위"),
    ("운송", "공급자 부담"), ("보증기간", "납품 후 12개월"), ("담당자", "______"), ("연락처", "______"),
    ("승인권자", "구매팀장"), ("단가 기준", "최근 견적서"), ("수량 단위", "EA"), ("도면", "최신 Rev. 적용"),
    ("긴급 발주", "대표 구두 승인 후 서면 보완"), ("변경 통보", "납기 3일 전까지"), ("분쟁", "당사 소재지 관할"),
    ("첨부", "견적서 사본"), ("비고란", "특이사항 기재"), ("문서 보관", "3년"), ("재발주", "동일 조건"),
    ("샘플 승인", "초도품 승인 후 양산"),
]
BODY = ["아래와 같이 발주합니다.", "본 발주서는 양사 합의된 거래 조건에 따릅니다.", "품명 | 규격 | 수량 | 단가 | 금액"]


def write(path, title, fields, body, note=None):
    d = docx.Document()
    d.add_heading(title, level=1)
    if note:
        d.add_paragraph(note)
    for k, v in fields:
        d.add_paragraph(f"{k}: {v}")
    for b in body:
        d.add_paragraph(b)
    d.save(path)


def main(seed=7):
    rng = random.Random(seed)
    (OUT / "changed").mkdir(parents=True, exist_ok=True)
    (OUT / "unchanged").mkdir(parents=True, exist_ok=True)

    old = list(BASE_FIELDS)
    new = list(BASE_FIELDS)
    gold = []
    keys = [k for k, _ in BASE_FIELDS]
    picked = rng.sample(keys, 16)
    # 기준값 변경 12개
    for k in picked[:12]:
        i = next(i for i, (kk, _) in enumerate(new) if kk == k)
        new[i] = (k, new[i][1] + " (변경)")
        gold.append({"field": k, "type": "기준값"})
    # 항목 삭제 4개
    for k in picked[12:16]:
        new = [(kk, v) for kk, v in new if kk != k]
        gold.append({"field": k, "type": "항목삭제"})
    # 항목 추가 4개
    for k, v in [("하자보수", "통보 후 7일 이내"), ("환경 기준", "RoHS 준수"), ("안전 교육", "출입 전 이수"), ("전자 결재", "그룹웨어 상신")]:
        new.append((k, v))
        gold.append({"field": k, "type": "항목추가"})

    write(OUT / "changed" / "발주서 양식.docx", "발 주 서", old, BODY)
    write(OUT / "changed" / "발주서 양식_개정.docx", "발 주 서", new, BODY, note="※ 2025.09 개정")
    write(OUT / "unchanged" / "발주서 양식.docx", "발 주 서", old, BODY)
    write(OUT / "unchanged" / "발주서 양식 - 복사본.docx", "발 주 서", old, BODY)

    (OUT / "gold.json").write_text(json.dumps({
        "description": "시드 평가 (개발자가 만든 정답). 제3자 블라인드 평가가 아님.",
        "seed": seed, "changed_fields": gold, "unchanged_pair_expected_changes": 0,
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{OUT}: 변경 {len(gold)}개 (기준값 12, 삭제 4, 추가 4)")


if __name__ == "__main__":
    main()
