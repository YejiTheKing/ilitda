"""parse_document: 파일에서 본문과 내부 정보를 꺼낸다.

파일 수정 날짜는 복사만 해도 바뀌므로 꺼내지 않는다.
버전 판별에는 문서 내용과 문서 안에 저장된 개정 정보만 쓴다.
"""
import re
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from ilitda.safety.paths import check_allowed


class UnreadableError(Exception):
    """읽을 수 없는 파일. 에이전트는 건너뛰고 사용자에게 알린다."""


def _pdf(path):
    import pypdf

    reader = pypdf.PdfReader(str(path))
    if reader.is_encrypted:
        raise UnreadableError("암호가 걸린 PDF")
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    info = reader.metadata or {}
    meta = {k.strip("/"): str(v) for k, v in info.items() if k in ("/CreationDate", "/ModDate")}
    return text, meta


def _docx(path):
    import docx

    doc = docx.Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            # 병합된 칸은 같은 내용이 반복되어 나오므로 이어진 중복을 뺀다
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if not cells or cells[-1] != t:
                    cells.append(t)
            lines.append(" | ".join(cells))
    props = doc.core_properties
    meta = {"revision": props.revision, "last_modified_by": props.last_modified_by}
    return "\n".join(lines), {k: v for k, v in meta.items() if v}


def _xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"[시트] {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                lines.append(" | ".join(cells).rstrip(" |"))
    return "\n".join(lines), {"sheets": [ws.title for ws in wb.worksheets]}


def _hwpx(path):
    # HWPX는 zip 안에 구역별 XML(Contents/section0.xml ...)이 들어 있다
    lines = []
    with zipfile.ZipFile(path) as z:
        sections = sorted(n for n in z.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n))
        if not sections:
            raise UnreadableError("HWPX 본문을 찾지 못함")
        for name in sections:
            root = ElementTree.fromstring(z.read(name))
            for el in root.iter():
                if el.tag.endswith("}p"):
                    text = "".join(t.text or "" for t in el.iter() if t.tag.endswith("}t"))
                    if text.strip():
                        lines.append(text)
    return "\n".join(lines), {}


# 한글 본문의 제어 문자. 인라인·확장 제어는 뒤에 7글자(14바이트)의 부가 정보가 따라온다
_HWP_SKIP8 = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}


def _hwp_chars(raw):
    out, i = [], 0
    while i < len(raw):
        code = ord(raw[i])
        if code in _HWP_SKIP8:
            i += 8
            continue
        if code in (10, 13):
            out.append("\n")
        elif code in (24, 30, 31):  # 하이픈·고정폭 빈칸
            out.append(" ")
        elif code >= 32:
            out.append(raw[i])
        i += 1
    text = re.sub(r"[ \t]+", " ", "".join(out))
    return "\n".join(ln.strip() for ln in text.splitlines() if ln.strip())


def _hwp(path):
    """구형 한글(.hwp 5.x). OLE 컨테이너 안의 본문 레코드에서 글자만 꺼낸다.

    표·그림 구조는 잃고 글자 순서만 남는다. 배포용(암호화) 문서는 읽지 못한다.
    """
    import struct
    import zlib

    import olefile

    if not olefile.isOleFile(str(path)):
        raise UnreadableError("한글 파일 구조가 아님 (손상되었거나 HWP 3.0 이하)")
    try:
        ole = olefile.OleFileIO(str(path))
    except Exception as e:   # 머리글만 남고 구조가 깨진 파일 (NotOleFileError 등)
        raise UnreadableError("손상된 한글 파일 (내용을 읽을 수 없음)") from e
    with ole:
        if not ole.exists("FileHeader"):
            raise UnreadableError("한글 파일 머리글이 없음")
        header = ole.openstream("FileHeader").read()
        flags = header[36]
        if flags & 0x02:
            raise UnreadableError("암호가 걸린 한글 파일")
        if flags & 0x04:
            raise UnreadableError("배포용 한글 파일 (본문이 잠겨 있음)")
        compressed = flags & 0x01
        sections = sorted((e for e in ole.listdir() if e[0] == "BodyText"), key=lambda e: int(e[1][7:]) if e[1][7:].isdigit() else 0)
        if not sections:
            raise UnreadableError("한글 본문을 찾지 못함")
        paragraphs = []  # 문단별 글자
        for entry in sections:
            data = ole.openstream(entry).read()
            if compressed:
                data = zlib.decompress(data, -15)
            i = 0
            while i + 4 <= len(data):
                head = struct.unpack_from("<I", data, i)[0]
                tag, size = head & 0x3FF, (head >> 20) & 0xFFF
                i += 4
                if size == 0xFFF:
                    size = struct.unpack_from("<I", data, i)[0]
                    i += 4
                if tag == 67:  # HWPTAG_PARA_TEXT
                    raw = data[i:i + size].decode("utf-16le", "ignore")
                    paragraphs.append(_hwp_chars(raw))
                i += size
    text = "\n".join(p for p in paragraphs if p)
    return text, {}


def _text(path):
    return Path(path).read_text(encoding="utf-8", errors="replace"), {}


PARSERS = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".hwpx": _hwpx, ".hwp": _hwp, ".txt": _text, ".md": _text}


def parse_document(path, allowed=None):
    """{"path", "name", "ext", "text", "meta"}를 돌려준다. 읽지 못하면 UnreadableError."""
    real = check_allowed(path, allowed)
    ext = real.suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise UnreadableError(f"지원하지 않는 형식: {ext or '확장자 없음'}")
    try:
        text, meta = parser(real)
    except UnreadableError:
        raise
    except Exception as e:  # 손상된 파일 등. 사용자 화면에 예외 이름을 그대로 내지 않는다
        raise UnreadableError("파일을 열지 못함 (손상되었을 수 있음)") from e

    text = unicodedata.normalize("NFC", text)
    if not text.strip():
        raise UnreadableError("글자를 찾지 못함 (스캔 이미지일 수 있음)")
    return {
        "path": str(real),
        "name": unicodedata.normalize("NFC", real.name),
        "ext": ext,
        "text": text,
        "meta": meta,
    }
